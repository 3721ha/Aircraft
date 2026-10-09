"""Replace only section 6.4 with the controlled component-ablation report."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Aircraft--v6.0.docx"
OUTPUT = ROOT / "xiaorong6.4.docx"
ABLATION_ROOT = ROOT / "results/final_protocol_20261007/component_ablation_10seed"
CONFLICT_ROWS = ROOT / "results/final_protocol_20261007/conflict_10seed/per_seed.json"


def fmt(value, digits=4):
    if value is None:
        return "N/A"
    return f"{float(value):.{digits}f}"


def set_run_font(run, name="宋体", size=10.5, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        name,
    )
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def style_paragraph(paragraph, size=10.5, align=None):
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size)


def add_text(doc, text, size=10.5, align=None):
    paragraph = doc.add_paragraph()
    paragraph.add_run(text)
    style_paragraph(paragraph, size=size, align=align)
    return paragraph


def add_heading(doc, text):
    paragraph = doc.add_paragraph(style="Heading 3")
    paragraph.add_run(text)
    style_paragraph(paragraph, size=11, align=None)
    for run in paragraph.runs:
        set_run_font(run, name="黑体", size=11, bold=True)
    return paragraph


def add_formula(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    set_run_font(run, name="Cambria Math", size=10.5)
    return paragraph


def add_table(doc, headers, rows, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, name="黑体", size=font_size, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = str(text)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, name="宋体", size=font_size)
    return table


def move_before(element, anchor):
    anchor.addprevious(element)


def find_heading(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph._p
    raise RuntimeError(f"heading not found: {prefix}")


def remove_between(start, end):
    current = start.getnext()
    while current is not None and current is not end:
        following = current.getnext()
        current.getparent().remove(current)
        current = following


def load_rows():
    summary = json.loads((ABLATION_ROOT / "statistical_summary/summary.json").read_text(encoding="utf-8"))
    deltas = json.loads((ABLATION_ROOT / "statistical_summary/paired_deltas.json").read_text(encoding="utf-8"))
    conflict = json.loads(CONFLICT_ROWS.read_text(encoding="utf-8"))
    return summary, deltas, conflict


def summary_row(summary, variant):
    return next(row for row in summary if row["variant"] == variant)


def delta_row(deltas, component, metric):
    return next(row for row in deltas if row["component"] == component and row["metric"] == metric)


def main():
    summary, deltas, conflict = load_rows()
    shutil.copy2(SOURCE, OUTPUT)
    document = Document(OUTPUT)
    start = find_heading(document, "6.4 冲突仲裁消融")
    end = find_heading(document, "6.5 高冲突压力测试与能力边界")
    remove_between(start, end)

    # Add content at the document end, then move each element into the 6.4
    # position. This preserves all sections, figures, and tables outside 6.4.
    elements = []
    def keep(element):
        elements.append(element._p if hasattr(element, "_p") else element._tbl)

    keep(add_heading(document, "6.4.1 实验目的与整体配置对比"))
    keep(add_text(document, "第 6.4 节分两层回答两个问题。第一层比较无屏蔽、局部屏蔽、联合启发式和完整 DG-QP 四种安全层，考察安全链路的综合收益；第二层在规则覆盖、belief-STL 报告、名义策略、场景、随机种子和时间窗口完全一致的条件下，只关闭一个核心组件，估计依赖图、动态优先级和连续残差 QP 的独立净收益。前一层延续原实验协议，后一层是新增的受控组件消融。"))
    keep(add_text(document, "整体配置实验仍使用八类冲突模板、10 个种子（101、202、303、404、505、606、707、808、909、1001），每个种子每类模板运行 1 个 episode，并仅保留初始真值硬安全的样本。四种配置共用固定的 rule-agnostic 名义策略；差异仅来自安全层。表 6-3 汇总 80 个种子—场景单元的均值。"))

    conflict_methods = ["NoShield", "PartialShield", "JointHeuristic", "DG-QP"]
    conflict_rows = []
    for method in conflict_methods:
        values = [row for row in conflict if row["method"] == method]
        conflict_rows.append([
            method,
            fmt(sum(row["truth_hard_violation_rate"] for row in values) / len(values)),
            fmt(sum(row["target_hard_success_rate"] for row in values) / len(values)),
            fmt(sum(row["intervention_rate"] for row in values) / len(values)),
            fmt(sum(row["qp_infeasible_rate"] for row in values) / len(values)),
        ])
    table = add_table(document,
        ["安全层配置", "真值硬违规率↓", "目标规则成功率↑", "干预/步", "QP回退率↓"],
        conflict_rows,
    )
    keep(table)
    keep(add_text(document, "表 6-3  八类冲突模板的整体安全层对比（80 个种子—场景单元）", size=9, align=WD_ALIGN_PARAGRAPH.CENTER))
    keep(add_text(document, "整体配置结果表明，DG-QP 将平均真值硬违规率降至 0.0483，目标规则成功率提高到 0.7625，但干预频率和 QP 回退率也最高。该结果支持完整安全链路的综合作用，不足以把收益单独归因于某一个内部组件。"))

    keep(add_heading(document, "6.4.2 受控组件消融设计"))
    keep(add_text(document, "受控消融固定完整规则覆盖集合，即所有变体都计算同一组 S/I/F/C/M/E 规则的 belief-STL 报告，并使用相同的硬规则阈值、置信裕度、安全缓冲、动作候选集合、名义策略和求解时间窗口。唯一变化是安全层内部组件："))
    table = add_table(document,
        ["配置", "依赖图", "动态优先级", "连续残差 QP", "规则覆盖"],
        [
            ["DG-QP（完整）", "✓", "✓", "✓", "完整规则集合"],
            ["No-Graph", "✗", "✓", "✓", "同一完整规则集合"],
            ["Fixed-Priority", "✓", "固定 P0→P5", "✓", "同一完整规则集合"],
            ["Gate-Only", "✓", "✓", "✗（仅离散门控）", "同一完整规则集合"],
        ],
    )
    keep(table)
    keep(add_text(document, "表 6-4  受控组件消融配置。No-Graph 仍保留完整规则报告，但移除显式规则—智能体依赖边和规则—规则冲突事件；Fixed-Priority 使用固定 P0＞P1＞P2＞P3＞P4＞P5 顺序，并以规则编号作为同层次的确定性平局规则；Gate-Only 保留图和优先级门控，但不求解连续控制残差。" , size=9))
    keep(add_text(document, "连续控制修正以执行动作与名义动作的残差为代价，并对硬规则使用不可松弛约束。其统一形式写为："))
    keep(add_formula(document, "J(δu, ξ) = 1/2 ||δu||₂² + λ Σᵣ ξᵣ²"))
    keep(add_formula(document, "subject to  gᵣ(x̂ₜ₊₁, uₜ + δu) + ξᵣ ≥ κσᵣ + bᵣ,    ξᵣ = 0  (r ∈ hard rules)"))
    keep(add_text(document, "其中 δu 是连续控制残差，ξᵣ 是仅对可松弛规则开放的松弛量，σᵣ 是 belief 不确定性尺度，κ 是保守系数，bᵣ 是安全缓冲。DG-QP 变体使用 SLSQP 求解这一非线性约束残差问题；本文沿用 DG-QP 的历史命名，但不将其表述为具有全局最优性保证的严格二次规划。"))
    keep(add_text(document, "对每个种子，先在八类模板上求均值，再以种子作为配对统计单位。这样可以避免把同一种子内的多个场景误当作相互独立样本。组件 c 的方向统一净收益定义为："))
    keep(add_formula(document, "Δᵤ(c, m) = Mᵤ(DG-QP, m) − Mᵤ(Ablation_c, m)"))
    keep(add_text(document, "对于真值违规率和 QP 回退率等越低越好的指标，将上式的符号反向后再报告为 Δᵤ；因此表 6-6 中的正值表示完整 DG-QP 在该指标方向上更好。Reward、干预频率和在线耗时不被压缩为人为加权的单一总分，而是分别报告任务收益、干预代价和计算代价。"))

    keep(add_heading(document, "6.4.3 组件净收益结果"))
    controlled_rows = []
    for variant in ("DG-QP", "No-Graph", "Fixed-Priority", "Gate-Only"):
        row = summary_row(summary, variant)
        controlled_rows.append([
            variant,
            f"{fmt(row['mean_reward_mean'])} ± {fmt(row['mean_reward_ci95'])}",
            f"{fmt(row['joint_satisfaction_rate_mean'])} ± {fmt(row['joint_satisfaction_rate_ci95'])}",
            f"{fmt(row['post_truth_hard_violation_rate_mean'])} ± {fmt(row['post_truth_hard_violation_rate_ci95'])}",
            f"{fmt(row['target_hard_success_rate_mean'])} ± {fmt(row['target_hard_success_rate_ci95'])}",
            f"{fmt(row['mean_intervention_rate_mean'])} ± {fmt(row['mean_intervention_rate_ci95'])}",
            f"{fmt(row['mean_online_time_ms_mean'])} ± {fmt(row['mean_online_time_ms_ci95'])}",
        ])
    table = add_table(document,
        ["配置", "Reward", "Joint STL", "真值硬违规率", "目标成功率", "干预/步", "在线 ms"],
        controlled_rows,
        font_size=7.8,
    )
    keep(table)
    keep(add_text(document, "表 6-5  规则覆盖一致的组件消融总体结果（10 个种子，均值 ± 95% t 置信区间）", size=9, align=WD_ALIGN_PARAGRAPH.CENTER))

    selected = [
        ("dependency_graph", "rule_rule_conflict_step_rate", "依赖图", "规则冲突步率"),
        ("dependency_graph", "rule_rule_resolution_rate", "依赖图", "规则冲突解决率"),
        ("dependency_graph", "post_truth_hard_violation_rate", "依赖图", "真值硬违规率"),
        ("dynamic_priority", "joint_satisfaction_rate", "动态优先级", "Joint STL"),
        ("dynamic_priority", "post_truth_hard_violation_rate", "动态优先级", "真值硬违规率"),
        ("continuous_qp", "joint_satisfaction_rate", "连续 QP", "Joint STL"),
        ("continuous_qp", "post_truth_hard_violation_rate", "连续 QP", "真值硬违规率"),
        ("continuous_qp", "mean_reward", "连续 QP", "Reward"),
        ("continuous_qp", "mean_intervention_rate", "连续 QP", "干预/步"),
        ("continuous_qp", "mean_online_time_ms", "连续 QP", "在线 ms"),
    ]
    delta_rows = []
    for component, metric, label, metric_label in selected:
        row = delta_row(deltas, component, metric)
        delta_rows.append([
            label,
            metric_label,
            str(row["n"]),
            fmt(row["mean_delta_better"]),
            fmt(row["ci95_halfwidth"]),
            fmt(row["sign_p_two_sided"]),
        ])
    table = add_table(document,
        ["组件", "指标", "N", "方向净收益", "95% CI 半宽", "符号检验 p"],
        delta_rows,
        font_size=7.8,
    )
    keep(table)
    keep(add_text(document, "表 6-6  组件配对净收益。安全指标的正值表示完整 DG-QP 更优；干预/步和在线 ms 保留原始差值，正值分别表示完整配置干预更多或耗时更高。" , size=9, align=WD_ALIGN_PARAGRAPH.CENTER))

    keep(add_text(document, "受控消融给出三点需要区分的结论。第一，依赖图关闭后，Joint STL、真值硬违规率、目标规则成功率和 Reward 与完整配置完全一致，但规则冲突步率由 0.5175 降为 0，规则冲突解决率由 0.8295 降为 0。这说明依赖图在当前八类模板中的直接净收益首先体现为规则—智能体冲突的显式定位、记录和解决状态可追溯性；本批模板尚未产生额外的主安全指标差异。"))
    keep(add_text(document, "第二，Fixed-Priority 与动态优先级在十个种子的主要安全和任务指标上完全一致，说明当前模板中同时竞争同一离散动作的多规则提案较少，动态优先级没有被充分激活。该结果不能证明动态优先级无效，只能说明在本实验覆盖范围内尚未观察到其独立性能增益；后续应增加多规则同机同时触发且优先级排序会改变动作选择的定向场景。"))
    keep(add_text(document, "第三，去除连续 QP 后，Gate-Only 的 Joint STL 从 0.5625 降至 0.5500，真值硬违规率从 0.0483 升至 0.0504；对应方向净收益分别为 0.0125 和 0.0021，但十种子配对检验均未达到显著性。完整 DG-QP 的平均在线时间比 Gate-Only 高约 11.63 ms，平均每步干预多 0.0333 次，Reward 几乎不变。因而连续 QP 的证据应表述为“在当前协议下提供小幅安全改进并承担可测计算代价”，不能表述为在所有场景中显著优于门控。"))
    keep(add_text(document, "综合来看，原四配置实验支持完整安全链路的总体收益，新增受控消融则揭示了收益的组成：依赖图主要提供冲突语义的可追溯性，连续 QP 提供有限的连续动作安全修正，动态优先级的独立贡献需要更有针对性的多规则竞争场景才能充分检验。本节因此不把三个组件都宣称为已被主指标显著证明，而是同时报告有效性、代价和当前实验的识别边界。所有逐场景逐种子原始记录、统计表和配置清单保存在 `results/final_protocol_20261007/component_ablation_10seed/`。"))

    for element in elements:
        move_before(element, end)
    document.save(OUTPUT)
    print(json.dumps({"output": str(OUTPUT), "inserted_elements": len(elements)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
